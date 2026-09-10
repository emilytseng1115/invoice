from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


@dataclass(frozen=True)
class WordMatch:
    paragraph_index: int
    start: int
    end: int
    location: str
    preview: str


def _iter_table_paragraphs(table: Table, table_number: int):
    seen_cells: set[int] = set()
    for row_number, row in enumerate(table.rows, start=1):
        for column_number, cell in enumerate(row.cells, start=1):
            cell_id = id(cell._tc)
            if cell_id in seen_cells:
                continue
            seen_cells.add(cell_id)
            location = f"表格 {table_number}／第 {row_number} 列／第 {column_number} 欄"
            for paragraph in cell.paragraphs:
                yield paragraph, location
            for nested_number, nested_table in enumerate(cell.tables, start=1):
                yield from _iter_table_paragraphs(nested_table, nested_number)


def _document_paragraphs(document: DocumentObject) -> list[tuple[Paragraph, str]]:
    result: list[tuple[Paragraph, str]] = []
    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        result.append((paragraph, f"本文第 {paragraph_number} 段"))
    for table_number, table in enumerate(document.tables, start=1):
        result.extend(_iter_table_paragraphs(table, table_number))
    for section_number, section in enumerate(document.sections, start=1):
        for area_name, area in (("頁首", section.header), ("頁尾", section.footer)):
            for paragraph_number, paragraph in enumerate(area.paragraphs, start=1):
                result.append((paragraph, f"第 {section_number} 節{area_name}第 {paragraph_number} 段"))
    return result


def find_word_matches(docx_bytes: bytes, anchor_text: str) -> list[WordMatch]:
    anchor = anchor_text.strip()
    if not anchor:
        return []
    document = Document(BytesIO(docx_bytes))
    matches: list[WordMatch] = []
    for paragraph_index, (paragraph, location) in enumerate(_document_paragraphs(document)):
        text = paragraph.text
        start = 0
        while True:
            found = text.find(anchor, start)
            if found < 0:
                break
            matches.append(
                WordMatch(
                    paragraph_index=paragraph_index,
                    start=found,
                    end=found + len(anchor),
                    location=location,
                    preview=text.strip() or anchor,
                )
            )
            start = found + len(anchor)
    return matches


def _replace_span(paragraph: Paragraph, start: int, end: int, replacement: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(replacement)
        return

    positions: list[tuple[int, int]] = []
    cursor = 0
    for run in runs:
        positions.append((cursor, cursor + len(run.text)))
        cursor += len(run.text)

    start_index = len(runs) - 1
    end_index = len(runs) - 1
    for index, (left, right) in enumerate(positions):
        if left <= start <= right:
            start_index = index
            break
    for index, (left, right) in enumerate(positions):
        if left <= end <= right:
            end_index = index
            break

    start_offset = start - positions[start_index][0]
    end_offset = end - positions[end_index][0]
    if start_index == end_index:
        original = runs[start_index].text
        runs[start_index].text = original[:start_offset] + replacement + original[end_offset:]
        return

    prefix = runs[start_index].text[:start_offset]
    suffix = runs[end_index].text[end_offset:]
    runs[start_index].text = prefix + replacement
    for index in range(start_index + 1, end_index):
        runs[index].text = ""
    runs[end_index].text = suffix


def edit_docx(
    docx_bytes: bytes,
    matches: list[WordMatch],
    new_text: str,
    mode: str = "append",
    replace_text: str = "",
) -> bytes:
    if not matches:
        raise ValueError("沒有可修改的 Word 欄位。")
    document = Document(BytesIO(docx_bytes))
    paragraphs = _document_paragraphs(document)

    # Edit from the back so multiple matches in one paragraph retain their offsets.
    for match in sorted(matches, key=lambda item: (item.paragraph_index, item.start), reverse=True):
        paragraph = paragraphs[match.paragraph_index][0]
        if mode == "append":
            spacer = "" if new_text.startswith((" ", "\t")) else " "
            _replace_span(paragraph, match.end, match.end, spacer + new_text)
            continue

        if not replace_text:
            raise ValueError("請輸入要取代的原有文字。")
        found = paragraph.text.find(replace_text, match.end)
        target_paragraph = paragraph
        if found < 0:
            # Word forms often store the label and value in adjacent table cells.
            for next_index in range(match.paragraph_index + 1, min(match.paragraph_index + 6, len(paragraphs))):
                candidate = paragraphs[next_index][0]
                found = candidate.text.find(replace_text)
                if found >= 0:
                    target_paragraph = candidate
                    break
        if found < 0:
            raise ValueError(f"在欄位後方找不到要取代的文字：{replace_text}")
        _replace_span(target_paragraph, found, found + len(replace_text), new_text)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
